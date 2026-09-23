import math
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree.ElementTree import Element, SubElement, tostring

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patheffects as pe
from matplotlib.font_manager import FontProperties, findfont
from matplotlib.patches import FancyBboxPatch
from openpyxl import load_workbook
import trimesh

COASTER_SIZE = 100.0
COASTER_THICKNESS = 4.0
CORNER_RADIUS = 7.0
ART_DEPTH = 1.0
NAME_FONT = 'Roboto Condensed:style=Bold'
NUMBER_FONT = 'DejaVu Sans Condensed:style=Bold'

# Locked V4 style approved from JAMES 7
NAME_GAP = 0.55
NAME_OUTER = 1.20
NUM_GAP = 0.70
NUM_OUTER = 1.55
NAME_CENTER_Y = 37.0
NAME_END_DROP = 7.0
NAME_TARGET_WIDTH = 68.0
NAME_FONT_SIZE = 14.0
NAME_XSCALE = 0.78
NUMBER_CENTER_Y = -10.5
NUMBER_SIZE_SINGLE = 44.0
NUMBER_SIZE_DOUBLE = 42.0
NUMBER_XSCALE = 0.73


def safe_name(s):
    return ''.join(c if c.isalnum() else '_' for c in str(s)).strip('_')


def name_layout(name):
    name = str(name).upper().strip()
    n = len(name)
    if n <= 5:
        width = min(NAME_TARGET_WIDTH, 10.0 * (n - 1) if n > 1 else 0)
        fsize = NAME_FONT_SIZE
        xs = NAME_XSCALE
    elif n <= 8:
        width = NAME_TARGET_WIDTH
        fsize = 13.2
        xs = 0.75
    elif n <= 10:
        width = 72.0
        fsize = 11.8
        xs = 0.71
    else:
        width = 75.0
        fsize = 10.6
        xs = 0.67
    if n == 1:
        xs_pos = [0.0]
    else:
        xs_pos = [-width / 2 + i * width / (n - 1) for i in range(n)]
    out = []
    half = max(width / 2, 1.0)
    for ch, x in zip(name, xs_pos):
        y = NAME_CENTER_Y - NAME_END_DROP * (x / half) ** 2
        dydx = -2 * NAME_END_DROP * x / (half ** 2)
        rot = math.degrees(math.atan(dydx))
        rot = max(-30.0, min(30.0, rot))
        out.append((ch, x, y, rot, fsize, xs))
    return out


def scad_text_shape(text, size, font, xscale=1.0):
    esc = text.replace('\\', '\\\\').replace('"', '\\"')
    return f'scale([{xscale:.5f},1,1]) text("{esc}", size={size:.5f}, font="{font}", halign="center", valign="center", $fn=48);'


def ring_2d(inner_expr, gap, outer):
    return f'difference(){{ offset(delta={outer:.4f}) {{ {inner_expr} }} offset(delta={gap:.4f}) {{ {inner_expr} }} }}'


def name_2d(name):
    parts = []
    for ch, x, y, rot, fs, xs in name_layout(name):
        base = scad_text_shape(ch, fs, NAME_FONT, xs)
        transform = f'translate([{x:.5f},{y:.5f}]) rotate({rot:.5f}) {{ {base} }}'
        ring = ring_2d(transform, NAME_GAP, NAME_OUTER)
        parts.append(f'union(){{ {transform} {ring} }}')
    return 'union(){' + ' '.join(parts) + '}'


def number_2d(number):
    txt = str(number).upper().strip()
    if txt.isdigit():
        fs = NUMBER_SIZE_SINGLE if len(txt) == 1 else NUMBER_SIZE_DOUBLE
        xs = NUMBER_XSCALE if len(txt) == 1 else 0.70
    else:
        fs = 31.0
        xs = 0.70
    base = scad_text_shape(txt, fs, NUMBER_FONT, xs)
    transform = f'translate([0,{NUMBER_CENTER_Y:.5f}]) {{ {base} }}'
    ring = ring_2d(transform, NUM_GAP, NUM_OUTER)
    return f'union(){{ {transform} {ring} }}'


def artwork_2d(name, number, mirror_for_bottom=True):
    art = f'union(){{ {name_2d(name)} {number_2d(number)} }}'
    if mirror_for_bottom:
        return f'mirror([1,0,0]) {{ {art} }}'
    return art


def make_scad(name, number, part, face_down=True):
    art = artwork_2d(name, number, mirror_for_bottom=face_down)
    rounded = f'offset(r={CORNER_RADIUS}) square([{COASTER_SIZE-2*CORNER_RADIUS},{COASTER_SIZE-2*CORNER_RADIUS}], center=true);'
    if face_down:
        art3d = f'linear_extrude(height={ART_DEPTH}) {{ {art} }}'
    else:
        art3d = f'translate([0,0,{COASTER_THICKNESS-ART_DEPTH}]) linear_extrude(height={ART_DEPTH}) {{ {art} }}'
    full = f'linear_extrude(height={COASTER_THICKNESS}) {{ {rounded} }}'
    if part == 'black':
        geom = art3d
    elif part == 'white':
        geom = f'difference(){{ {full} {art3d} }}'
    else:
        raise ValueError(part)
    return f'$fn=80;\n{geom}\n'


def _candidate_openscad_paths():
    candidates = []
    if getattr(sys, 'frozen', False):
        base = Path(getattr(sys, '_MEIPASS', Path(sys.executable).parent))
        candidates += [
            base / 'openscad' / 'openscad.exe',
            Path(sys.executable).parent / 'openscad' / 'openscad.exe',
        ]
    env = os.environ.get('OPENSCAD_PATH')
    if env:
        candidates.append(Path(env))
    candidates += [
        Path(r'C:\Program Files\OpenSCAD\openscad.exe'),
        Path(r'C:\Program Files (x86)\OpenSCAD\openscad.exe'),
    ]
    which = shutil.which('openscad') or shutil.which('openscad.exe')
    if which:
        candidates.append(Path(which))
    return candidates


def get_openscad_exe():
    for p in _candidate_openscad_paths():
        if p.exists():
            return str(p)
    raise RuntimeError('OpenSCAD runtime was not found. Use the GitHub-built portable package, or install OpenSCAD.')


def run_openscad(scad_text, out_stl, workdir):
    scad_path = Path(workdir) / ('tmp_' + safe_name(Path(out_stl).stem) + '.scad')
    scad_path.write_text(scad_text, encoding='utf-8')
    exe = get_openscad_exe()
    p = subprocess.run([exe, '-o', str(out_stl), str(scad_path)], capture_output=True, text=True, timeout=180)
    if p.returncode != 0:
        raise RuntimeError('OpenSCAD failed:\n' + p.stderr[-4000:])
    if not Path(out_stl).exists() or Path(out_stl).stat().st_size < 100:
        raise RuntimeError('OpenSCAD produced no usable STL')


def mesh_to_xml(mesh, obj_id, name, pid, pindex):
    obj = Element('object', {'id': str(obj_id), 'type': 'model', 'name': name, 'pid': str(pid), 'pindex': str(pindex)})
    mesh_el = SubElement(obj, 'mesh')
    verts = SubElement(mesh_el, 'vertices')
    for v in mesh.vertices:
        SubElement(verts, 'vertex', {'x': f'{v[0]:.5f}', 'y': f'{v[1]:.5f}', 'z': f'{v[2]:.5f}'})
    tris = SubElement(mesh_el, 'triangles')
    for f in mesh.faces:
        SubElement(tris, 'triangle', {'v1': str(int(f[0])), 'v2': str(int(f[1])), 'v3': str(int(f[2]))})
    return obj


def write_3mf(white_mesh, black_mesh, out_path, model_name):
    ns = 'http://schemas.microsoft.com/3dmanufacturing/core/2015/02'
    mns = 'http://schemas.microsoft.com/3dmanufacturing/material/2015/02'
    model = Element('model', {'unit': 'millimeter', 'xml:lang': 'en-US', 'xmlns': ns, 'xmlns:m': mns})
    resources = SubElement(model, 'resources')
    basemat = SubElement(resources, 'm:basematerials', {'id': '10'})
    SubElement(basemat, 'm:base', {'name': 'White', 'displaycolor': '#FFFFFF'})
    SubElement(basemat, 'm:base', {'name': 'Black', 'displaycolor': '#000000'})
    resources.append(mesh_to_xml(white_mesh, 1, 'WHITE_BODY', 10, 0))
    resources.append(mesh_to_xml(black_mesh, 2, 'BLACK_ART', 10, 1))
    comp = SubElement(resources, 'object', {'id': '3', 'type': 'model', 'name': model_name})
    comps = SubElement(comp, 'components')
    SubElement(comps, 'component', {'objectid': '1'})
    SubElement(comps, 'component', {'objectid': '2'})
    build = SubElement(model, 'build')
    SubElement(build, 'item', {'objectid': '3'})
    xml = b'<?xml version="1.0" encoding="UTF-8"?>' + tostring(model, encoding='utf-8')
    content_types = '''<?xml version="1.0" encoding="UTF-8"?>\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="model" ContentType="application/vnd.ms-package.3dmanufacturing-3dmodel+xml"/></Types>'''
    rels = '''<?xml version="1.0" encoding="UTF-8"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Target="/3D/3dmodel.model" Id="rel0" Type="http://schemas.microsoft.com/3dmanufacturing/2013/01/3dmodel"/></Relationships>'''
    with zipfile.ZipFile(out_path, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', content_types)
        z.writestr('_rels/.rels', rels)
        z.writestr('3D/3dmodel.model', xml)


def validate(white, black):
    for m in (white, black):
        if m.bounds[0][0] < -50.01 or m.bounds[1][0] > 50.01 or m.bounds[0][1] < -50.01 or m.bounds[1][1] > 50.01:
            raise RuntimeError('Mesh exceeds 100 mm coaster boundary')
        if m.bounds[0][2] < -0.01 or m.bounds[1][2] > 4.01:
            raise RuntimeError('Mesh exceeds 4 mm thickness')
    if not white.is_watertight or not black.is_watertight:
        raise RuntimeError('One or more generated meshes are not watertight')
    return {
        'white_watertight': bool(white.is_watertight),
        'black_watertight': bool(black.is_watertight),
        'white_bounds': white.bounds.tolist(),
        'black_bounds': black.bounds.tolist(),
    }


def draw_finished_preview(name, number, path):
    fig, ax = plt.subplots(figsize=(7, 7), dpi=140)
    ax.set_aspect('equal')
    ax.axis('off')
    ax.set_xlim(-55, 55)
    ax.set_ylim(-55, 55)
    patch = FancyBboxPatch((-50, -50), 100, 100, boxstyle='round,pad=0,rounding_size=7', facecolor='#f5f5f3', edgecolor='#ddddda', linewidth=1.0)
    ax.add_patch(patch)
    fp_name = FontProperties(fname=findfont(FontProperties(family='Roboto Condensed', weight='bold')))
    fp_num = FontProperties(fname=findfont(FontProperties(family='DejaVu Sans', weight='bold', stretch='condensed')))
    for ch, x, y, rot, fs, _xs in name_layout(name):
        t = ax.text(x, y, ch, ha='center', va='center', rotation=rot, fontproperties=fp_name, fontsize=fs * 3.0, color='black')
        t.set_path_effects([pe.Stroke(linewidth=2.3, foreground='white'), pe.Stroke(linewidth=3.5, foreground='black'), pe.Normal()])
    txt = str(number).upper().strip()
    fs = (NUMBER_SIZE_SINGLE if txt.isdigit() and len(txt) == 1 else NUMBER_SIZE_DOUBLE if txt.isdigit() else 31.0) * 1.35
    t = ax.text(0, NUMBER_CENTER_Y, txt, ha='center', va='center', fontproperties=fp_num, fontsize=fs, color='black')
    t.set_path_effects([pe.Stroke(linewidth=5.5, foreground='black'), pe.Stroke(linewidth=3.2, foreground='white'), pe.Normal()])
    fig.savefig(path, bbox_inches='tight', pad_inches=0.05)
    plt.close(fig)


def generate(name, number, outdir, face_down=True, prefix=None, keep_stl=True, make_preview=True):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    stem = prefix or f'{safe_name(name)}_{safe_name(number)}'
    with tempfile.TemporaryDirectory(prefix='rlo3mf_') as td:
        wstl = Path(td) / 'WHITE_BODY.stl'
        bstl = Path(td) / 'BLACK_ART.stl'
        run_openscad(make_scad(name, number, 'white', face_down), wstl, td)
        run_openscad(make_scad(name, number, 'black', face_down), bstl, td)
        white = trimesh.load_mesh(wstl, process=True)
        black = trimesh.load_mesh(bstl, process=True)
        checks = validate(white, black)
        suffix = 'FACE_DOWN' if face_down else 'FACE_UP'
        out3mf = outdir / f'{stem}_{suffix}.3mf'
        write_3mf(white, black, out3mf, f'{str(name).upper()}_{str(number).upper()}')
        if keep_stl:
            shutil.copy2(wstl, outdir / f'{stem}_WHITE_BODY.stl')
            shutil.copy2(bstl, outdir / f'{stem}_BLACK_ART.stl')
    prev = None
    if make_preview:
        prev = outdir / f'{stem}_finished_preview.png'
        draw_finished_preview(name, number, prev)
    return out3mf, checks, prev


def read_players_xlsx(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(v).strip().lower() if v is not None else '' for v in rows[0]]
    def idx_for(options, fallback):
        for opt in options:
            if opt in headers:
                return headers.index(opt)
        return fallback
    name_i = idx_for(['name', 'player', 'player name', 'surname'], 0)
    num_i = idx_for(['number', 'no', 'no.', 'shirt number'], 2)
    pos_i = idx_for(['position', 'pos'], 1)
    players = []
    for row in rows[1:]:
        if not row or name_i >= len(row) or not row[name_i]:
            continue
        name = str(row[name_i]).strip()
        number = '' if num_i >= len(row) or row[num_i] is None else str(row[num_i]).strip()
        position = '' if pos_i >= len(row) or row[pos_i] is None else str(row[pos_i]).strip()
        if number.endswith('.0') and number[:-2].isdigit():
            number = number[:-2]
        if not number:
            continue
        players.append({'name': name, 'number': number, 'position': position})
    return players
